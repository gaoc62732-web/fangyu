# -*- coding: utf-8 -*-
"""将本审核包地图导出的记录 JSON，按政区名称和代码合并进手册 XLSX。"""
import argparse
import json
import os
import re
import sys
from collections import defaultdict

import openpyxl
from openpyxl.utils import column_index_from_string


def clean(value):
    return re.sub(r"\s+", "", str(value or "").replace("√", "").replace("●", ""))


def code(value):
    if value is None:
        return ""
    text = str(value).strip()
    return text[:-2] if text.endswith(".0") else text


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--records", help="地图导出记录 JSON")
    source.add_argument("--site", help="带 saved-state 的地图 HTML")
    parser.add_argument("--base", required=True, help="基础手册 XLSX")
    parser.add_argument("--out", required=True, help="输出 XLSX")
    args = parser.parse_args()
    if os.path.abspath(args.base) == os.path.abspath(args.out):
        parser.error("--out 不能覆盖 --base")
    sitemap_path = os.path.join(os.path.dirname(__file__), "sitemap.json")
    sitemap = json.load(open(sitemap_path, encoding="utf-8"))
    if args.records:
        record = json.load(open(args.records, encoding="utf-8"))
    else:
        html = open(args.site, encoding="utf-8").read()
        match = re.search(r'<script[^>]*id="saved-state"[^>]*>(.*?)</script>', html, re.S)
        if not match:
            sys.exit("未找到 saved-state，请从地图导出记录 JSON")
        record = json.loads(match[1])
    if record.get("format") != "fangyu-map-records" or not isinstance(record.get("changes"), dict):
        sys.exit("记录文件格式不正确")
    if record.get("fingerprint") != sitemap["fingerprint"]:
        sys.exit("记录指纹与本审核包地图不一致；请先在本包地图中导入记录，再重新导出")

    workbook = openpyxl.load_workbook(args.base)
    sheet = next((workbook[n] for n in workbook.sheetnames if "统计名录" in n), None)
    if sheet is None:
        sys.exit("基础手册缺少【统计名录】工作表")
    by_code, by_name = defaultdict(list), defaultdict(list)
    for row in range(2, sheet.max_row + 1):
        c = code(sheet.cell(row, 1).value)
        if not re.fullmatch(r"\d{6}", c):
            continue
        name = next((clean(sheet.cell(row, col).value) for col in (4, 3, 2) if sheet.cell(row, col).value), "")
        by_code[c].append(row)
        if name:
            by_name[name].append(row)

    def locate(meta):
        candidates = by_code.get(meta["code"], [])
        named = [r for r in candidates if clean(sheet.cell(r, 4).value or sheet.cell(r, 3).value or sheet.cell(r, 2).value) == clean(meta["name"])]
        if len(named) == 1:
            return named[0]
        same_name = [r for r in by_name.get(clean(meta["name"]), []) if code(sheet.cell(r, 1).value)[:2] == meta["code"][:2]]
        return same_name[0] if len(same_name) == 1 else None

    applied, unchanged, skipped = 0, 0, []
    for cell, value in record["changes"].items():
        match = re.fullmatch(r"([B-R])(\d+)", cell)
        if not match or not isinstance(value, str):
            skipped.append((cell, "单元格或文本格式不正确"))
            continue
        col, row = match[1], int(match[2])
        meta = sitemap["rowMeta"].get(str(row))
        if not meta:
            skipped.append((cell, "不在本版地图名录中"))
            continue
        if col in "BCD" and meta["regionCell"] != cell:
            skipped.append((cell, "该行不是此政区级别的名称格"))
            continue
        target_row = locate(meta)
        if target_row is None:
            skipped.append((cell, "在目标手册中无法唯一定位政区"))
            continue
        target = sheet.cell(target_row, column_index_from_string(col))
        if str(target.value or "").strip() == value.strip():
            unchanged += 1
            continue
        target.value = value
        applied += 1
    workbook.save(args.out)
    print(f"已保存：{args.out}")
    print(f"应用 {applied} 格、原值相同 {unchanged} 格、跳过 {len(skipped)} 格")
    for cell, reason in skipped:
        print(f"[跳过] {cell}：{reason}")


if __name__ == "__main__":
    main()
